
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.chart import (
    ScatterChart,
    Reference,
    Series,
)

wb = openpyxl.load_workbook()
ws = wb.active

chart = ScatterChart()
chart.title = "Scatter Chart"
chart.style = 13
chart.x_axis.title = 'Size'
chart.y_axis.title = 'Percentage'

xvalues = Reference(ws, min_col=1, min_row=2, max_row=7)
for i in range(2, 4):
    values = Reference(ws, min_col=i, min_row=1, max_row=7)
    series = Series(values, xvalues, title_from_data=True)
    chart.series.append(series)

ws.add_chart(chart, "A10")

wb.save("scatter.xlsx")

='[08_29_22_MM220815-1_Cond6_30 ccm wet H2_ 3 bar flat pressure block_0.6V_7.xlsx]08_29_22_MM220815-1_Cond6'!$A:$A,'[08_29_22_MM220815-1_Cond6_30 ccm wet H2_ 3 bar 
flat pressure block_0.6V_7.xlsx]08_29_22_MM220815-1_Cond6'!$C:$C